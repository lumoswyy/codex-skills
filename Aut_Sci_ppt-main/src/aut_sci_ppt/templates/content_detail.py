"""
详细内容页 - 终版溢出保护
"""
from .base import BaseTemplate, hex_to_rgb
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN
from lxml import etree
import re


def _set_lnsp(tf, pt_val):
    for para in tf.paragraphs:
        pPr = para._p.get_or_add_pPr()
        el = pPr.find(qn("a:lnSpc"))
        if el is None: el = etree.SubElement(pPr, qn("a:lnSpc"))
        sp = el.find(qn("a:spcPts"))
        if sp is None: sp = etree.SubElement(el, qn("a:spcPts"))
        sp.set("val", str(int(pt_val * 100)))


def _calc(texts, avail_h, cw):
    for fs in range(20, 9, -1):
        cpl = max(1, int((cw - 0.18) * 72 / (fs * 0.58)))
        heights = [fs / 72 * 1.28 * max(1, -(-len(str(t)) // cpl)) + 0.06 for t in texts]
        if sum(heights) <= avail_h:
            return fs, heights
    fs = 9
    cpl = max(1, int((cw - 0.18) * 72 / (fs * 0.58)))
    heights = [fs / 72 * 1.28 * max(1, -(-len(str(t)) // cpl)) + 0.06 for t in texts]
    return fs, heights


class ContentDetailTemplate(BaseTemplate):
    HEADER_Y = 0.68
    FOOT_Y   = 7.28
    SEC_H    = 0.40

    def render(self, data):
        self.draw_nav_sidebar()
        cx = self.CONTENT_X
        cw = self.W - cx - 0.22

        title    = getattr(data, "title",   "")
        points   = getattr(data, "points",  []) or []
        results  = getattr(data, "results", []) or []
        part_num = getattr(data, "part_num","")

        self.draw_header_bar(part_num=part_num, title=title)

        y       = self.HEADER_Y + 0.14
        avail_h = self.FOOT_Y - y
        has_res = bool(results)

        if has_res and points:
            pts_avail = avail_h * 0.52 - self.SEC_H - 0.10
            res_avail = avail_h * 0.42 - self.SEC_H - 0.10
        elif points:
            pts_avail = avail_h - self.SEC_H - 0.10
            res_avail = 0
        else:
            pts_avail = 0
            res_avail = avail_h - self.SEC_H - 0.10

        if points:
            y = self._hd(cx, cw, y, "■  研究要点", self.COLOR_PRIMARY)
            y = self._block(points, cx, cw, y, pts_avail, self.COLOR_PRIMARY, "➢  ")

        if has_res:
            y += 0.10
            y = self._hd(cx, cw, y, "◆  主要结论", self.COLOR_RED)
            self._block(results, cx, cw, y, res_avail, self.COLOR_RED, "◆  ")

    def _hd(self, cx, cw, y, text, bg):
        self.add_rect(cx, y, cw, self.SEC_H, fill_color=bg)
        self.add_rect(cx, y, 0.07, self.SEC_H, fill_color=self.COLOR_GOLD)
        self.add_textbox(cx + 0.16, y + 0.05, cw - 0.2, self.SEC_H - 0.08,
            text=text, font_name=self.FONT_CN, font_size=17,
            bold=True, color=self.COLOR_WHITE)
        return y + self.SEC_H + 0.10

    def _block(self, items, cx, cw, y, avail_h, bullet_color, bullet):
        texts   = [str(it).strip() for it in items]
        fs, heights = _calc(texts, avail_h, cw)
        total   = sum(heights)
        gap     = max(0, avail_h - total) / len(texts)
        line_sp = fs * 1.28
        for i, text in enumerate(texts):
            box_h = heights[i] + gap
            tb = self.slide.shapes.add_textbox(
                Inches(cx + 0.05), Inches(y),
                Inches(cw - 0.05), Inches(box_h))
            tf = tb.text_frame; tf.word_wrap = True
            _set_lnsp(tf, line_sp)
            p = tf.paragraphs[0]
            r0 = p.add_run(); r0.text = bullet
            r0.font.name = self.FONT_EN; r0.font.size = Pt(fs)
            r0.font.color.rgb = hex_to_rgb(bullet_color)
            self._render_rich(p, text, fs)
            y += box_h
        return y

    def _render_rich(self, paragraph, text, size=14):
        for part in re.split(r"(\*\*.*?\*\*)", text):
            if part.startswith("**") and part.endswith("**"):
                r = paragraph.add_run(); r.text = part[2:-2]
                r.font.name = self.FONT_EN; r.font.size = Pt(size)
                r.font.bold = True
                r.font.color.rgb = hex_to_rgb(self.COLOR_RED)
            elif part:
                r = paragraph.add_run(); r.text = part
                r.font.name = self.FONT_EN; r.font.size = Pt(size)
                r.font.color.rgb = hex_to_rgb(self.COLOR_BLACK)
