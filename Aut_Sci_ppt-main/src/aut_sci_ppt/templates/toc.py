"""
目录页 - 无导航栏版，全页宽布局，章节标题 >= 28pt
"""
from .base import BaseTemplate, hex_to_rgb
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

class TOCTemplate(BaseTemplate):
    def render(self, data):
        # ── 不调用 draw_nav_sidebar() ── 全页宽布局
        cx = 0.5    # 从左边 0.5 英寸开始
        cw = self.W - 1.0

        # 深蓝顶部页眉（全宽）
        self.add_rect(0, 0, self.W, 0.72, fill_color=self.COLOR_PRIMARY)
        self.add_rect(0, 0, 0.1, 0.72, fill_color=self.COLOR_GOLD)
        self.add_rect(self.W - 0.1, 0, 0.1, 0.72, fill_color=self.COLOR_GOLD)
        self.add_textbox(0.2, 0.1, self.W - 0.4, 0.52,
            text="目  录   CONTENTS", font_name=self.FONT_CN,
            font_size=26, bold=True, color=self.COLOR_WHITE,
            align=PP_ALIGN.CENTER)

        items = getattr(data, "items", [])
        if not items: return

        n = len(items)
        start_top = 0.88
        avail_h   = self.H - start_top - 0.25
        TITLE_SIZE = 28
        # 每项高度：保证 28pt 能放下，不足时等分
        MIN_H  = 0.72
        item_h = max(MIN_H, avail_h / max(n, 1))
        if item_h * n > avail_h:
            item_h = avail_h / n

        # 两列布局（章节数 > 4 时）
        use_two_col = (n > 4)
        if use_two_col:
            col_w  = cw / 2 - 0.2
            cols   = [cx, cx + cw / 2 + 0.1]
            rows   = -(-n // 2)   # ceil
            item_h = max(MIN_H, avail_h / max(rows, 1))
        else:
            col_w = cw
            cols  = [cx]
            rows  = n

        for i, item in enumerate(items):
            if use_two_col:
                col_idx = i % 2
                row_idx = i // 2
            else:
                col_idx = 0
                row_idx = i

            top  = start_top + row_idx * item_h
            x    = cols[col_idx]
            num  = item.num   if hasattr(item, "num")   else str(i + 1)
            title= item.title if hasattr(item, "title") else str(item)

            # 编号色块
            box_s = min(item_h - 0.12, 0.58)
            self.add_rect(x, top + 0.06, 0.6, box_s, fill_color=self.COLOR_PRIMARY)
            self.add_textbox(x, top + 0.06, 0.6, box_s,
                text=num, font_name=self.FONT_EN, font_size=20,
                bold=True, color=self.COLOR_WHITE, align=PP_ALIGN.CENTER)

            # 标题（>= 28pt）
            txBox = self.slide.shapes.add_textbox(
                Inches(x + 0.72), Inches(top + 0.1),
                Inches(col_w - 0.82), Inches(item_h - 0.15))
            tf = txBox.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = title
            run.font.name = self.FONT_CN
            run.font.size = Pt(TITLE_SIZE)
            run.font.bold = True
            run.font.color.rgb = hex_to_rgb(self.COLOR_PRIMARY)

            # 下划线分隔
            if (not use_two_col and i < n - 1) or \
               (use_two_col and row_idx < rows - 1):
                self.add_line(x, top + item_h - 0.04,
                              col_w, color="DDDDDD", width_pt=0.5)
